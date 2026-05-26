"""
Creates Bank Rec_2.xlsx with realistic sample data for testing reconcile.py
Sheets required: Bank Data, Sage Data, Unposted Bank, Extra In Sage
"""
import openpyxl
from openpyxl import Workbook
import datetime

wb = Workbook()

# --- Bank Data sheet ---
bank = wb.active
bank.title = "Bank Data"
bank.append(["Date", "Ref", "Type", "Account", "Description", "Paid Out", "Paid In"])

bank_rows = [
    # (date, ref, type, account, description, paid_out, paid_in)
    (datetime.datetime(2024, 1, 5), "CHQ001", "CHQ", "1001", "RENT PAYMENT JAN 2024",    5000.00, None),
    (datetime.datetime(2024, 1, 8), "TRF001", "TRF", "1001", "SALARY TRANSFER JAN",       None,   12000.00),
    (datetime.datetime(2024, 1, 10),"DD001",  "DD",  "1001", "ELECTRICITY BILL JAN",      350.00, None),
    (datetime.datetime(2024, 1, 12),"TRF002", "TRF", "1001", "CUSTOMER PAYMENT INV101",   None,   3500.00),
    (datetime.datetime(2024, 1, 15),"CHQ002", "CHQ", "1001", "SUPPLIER PAYMENT XYZ",      1200.00,None),
    # split payment: bank has one entry, sage has two entries that add up
    (datetime.datetime(2024, 1, 18),"TRF003", "TRF", "1001", "CUSTOMER PAYMENT INV102",   None,   2000.00),
    # unreconciled — no match in sage
    (datetime.datetime(2024, 1, 20),"DD002",  "DD",  "1001", "INTERNET SERVICE FEE",      99.00,  None),
]

for r in bank_rows:
    bank.append(list(r))

# --- Sage Data sheet ---
sage = wb.create_sheet("Sage Data")
sage.append(["Account", "Ref", "Type", "Date", "Description In", "Description Out", "Amount In", "Amount Out"])

sage_rows = [
    # (account, ref, type, date_str, desc_in, desc_out, amt_in, amt_out)
    ("1001", "CHQ001", "CHQ", "05/01/2024", None,                          "RENT PAYMENT JAN 2024",    None,    5000.00),
    ("1001", "TRF001", "TRF", "08/01/2024", "SALARY TRANSFER JAN",         None,                       12000.00,None),
    ("1001", "DD001",  "DD",  "10/01/2024", None,                          "ELECTRICITY BILL JAN",     None,    350.00),
    ("1001", "TRF002", "TRF", "12/01/2024", "CUSTOMER PAYMENT INV101",     None,                       3500.00, None),
    ("1001", "CHQ002", "CHQ", "15/01/2024", None,                          "SUPPLIER PAYMENT XYZ",     None,    1200.00),
    # two sage lines that together = bank TRF003 of 2000
    ("1001", "TRF003A","TRF", "18/01/2024", "CUSTOMER PAYMENT INV102",     None,                       1200.00, None),
    ("1001", "TRF003B","TRF", "18/01/2024", "CUSTOMER PAYMENT INV102",     None,                       800.00,  None),
]

for r in sage_rows:
    sage.append(list(r))

# --- Output sheets (empty, headers added by reconcile.py) ---
wb.create_sheet("Unposted Bank")
wb.create_sheet("Extra In Sage")

wb.save("Bank Rec_2.xlsx")
print("Sample file created: Bank Rec_2.xlsx")
print("Sheets:", wb.sheetnames)
