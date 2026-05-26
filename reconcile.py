import openpyxl
import datetime

filename = "Bank Rec_2.xlsx"
wb = openpyxl.load_workbook(filename)

reconciled_bank_row_list1 = []
unreconciled_bank_row_list1 = []
final_unreconciled_bank_list = []
cummulative_reconciled_bank_rows = []

bank_sheet = wb["Bank Data"]
sage_sheet = wb["Sage Data"]
unposted_bank_sheet = wb["Unposted Bank"]
extra_in_sage_sheet = wb["Extra In Sage"]

unposted_bank_row = 2
extra_in_sage_row = 2

# Copy headers to output sheets
for col in range(1, 8):
    unposted_bank_sheet.cell(1, column=col).value = bank_sheet.cell(1, column=col).value
for col in range(1, 9):
    extra_in_sage_sheet.cell(1, column=col).value = sage_sheet.cell(1, column=col).value


for bank_row in range(2, bank_sheet.max_row + 1):
    unreconciled_bank_row_list1.append(bank_row)
    bank_trans_desc_cell = bank_sheet.cell(bank_row, column=5).value
    bank_trans_split = bank_trans_desc_cell.split()
    bank_trans_join = bank_trans_split[0] + " " + bank_trans_split[1]
    bank_trans_desc_cell = bank_trans_join

    bank_date = bank_sheet.cell(bank_row, column=1).value
    date = datetime.datetime.date(bank_date)
    xdate = date.strftime("%d/%m/%Y")
    bank_date = datetime.datetime.strptime(xdate, "%d/%m/%Y").date()

    payment_type = ""
    bank_col_6 = bank_sheet.cell(bank_row, column=6).value
    bank_col_7 = bank_sheet.cell(bank_row, column=7).value

    if isinstance(bank_col_6, (float, int)) and bank_col_6 is not None and bank_col_6 > 0:
        bank_price_cell = bank_col_6
        payment_type = "paid out"
    else:
        bank_price_cell = bank_col_7
        payment_type = "paid in"

    for sage_row in range(2, sage_sheet.max_row + 1):
        sage_trans_desc_cell = sage_sheet.cell(sage_row, column=5 if payment_type == "paid in" else 6).value
        sage_date = sage_sheet.cell(sage_row, column=4).value
        sage_date = datetime.datetime.strptime(sage_date, "%d/%m/%Y").date()

        col_7 = sage_sheet.cell(sage_row, column=7).value
        col_8 = sage_sheet.cell(sage_row, column=8).value
        if isinstance(col_7, (float, int)) and col_7 is not None and col_7 > 0:
            sage_price_cell = col_7
        elif isinstance(col_8, (float, int)) and col_8 is not None and col_8 > 0:
            sage_price_cell = col_8

        unreconciled_bank_row_list1.append(bank_row)

        if bank_row in reconciled_bank_row_list1:
            pass
        elif bank_price_cell == sage_price_cell and bank_date == sage_date:
            reconciled_bank_row_list1.append(bank_row)
            break

        if bank_row not in reconciled_bank_row_list1:
            if sage_trans_desc_cell and bank_trans_desc_cell in sage_trans_desc_cell:
                if bank_price_cell == sage_price_cell and bank_date == sage_date:
                    reconciled_bank_row_list1.append(bank_row)

                for sage_row2 in range(2, sage_sheet.max_row + 1):
                    sage_trans_desc_cell = sage_sheet.cell(sage_row2, column=5 if payment_type == "paid in" else 6).value
                    sage_date2 = sage_sheet.cell(sage_row2, column=4).value
                    sage_date2 = datetime.datetime.strptime(sage_date2, "%d/%m/%Y").date()
                    col_7_2 = sage_sheet.cell(sage_row2, column=7).value
                    col_8_2 = sage_sheet.cell(sage_row2, column=8).value
                    if isinstance(col_7_2, (float, int)) and col_7_2 is not None and col_7_2 > 0:
                        sage_price_cell2 = col_7_2
                    elif isinstance(col_8_2, (float, int)) and col_8_2 is not None and col_8_2 > 0:
                        sage_price_cell2 = col_8_2

                    if bank_price_cell == sage_price_cell2 and bank_date == sage_date2:
                        reconciled_bank_row_list1.append(bank_row)
                        break
                    elif bank_price_cell == sage_price_cell2 and bank_date != sage_date2:
                        unreconciled_bank_row_list1.append(bank_row)

                if bank_price_cell > sage_price_cell:
                    total_sage_price = 0
                    for row3 in range(2, sage_sheet.max_row + 1):
                        sage_trans_desc_cell2 = sage_sheet.cell(row3, column=5 if payment_type == "paid in" else 6).value
                        if sage_trans_desc_cell2 and bank_trans_desc_cell in sage_trans_desc_cell2:
                            col_7_3 = sage_sheet.cell(row3, column=7).value
                            col_8_3 = sage_sheet.cell(row3, column=8).value
                            if isinstance(col_7_3, (float, int)) and col_7_3 is not None and col_7_3 > 0:
                                total_sage_price += round(col_7_3, 2)
                            elif isinstance(col_8_3, (float, int)) and col_8_3 is not None and col_8_3 > 0:
                                total_sage_price += round(col_8_3, 2)
                        total_sage_price = round(total_sage_price, 2)

                    if bank_price_cell == total_sage_price:
                        reconciled_bank_row_list1.append(bank_row)
                    elif bank_price_cell < total_sage_price:
                        unreconciled_bank_row_list1.append(bank_row)
                        total_bank_price = 0
                        temp_cummulative_reconciled_bank_rows = []
                        for row4 in range(2, bank_sheet.max_row + 1):
                            bank_trans_desc_cell2 = bank_sheet.cell(row4, column=5).value
                            bank_trans_split2 = bank_trans_desc_cell2.split()
                            bank_trans_desc_cell2 = bank_trans_split2[0] + " " + bank_trans_split2[1]
                            if bank_trans_desc_cell == bank_trans_desc_cell2:
                                temp_cummulative_reconciled_bank_rows.append(row4)
                                bank_col_6 = bank_sheet.cell(row4, column=6).value
                                bank_col_7 = bank_sheet.cell(row4, column=7).value
                                if isinstance(bank_col_6, (float, int)) and bank_col_6 is not None and bank_col_6 > 0:
                                    total_bank_price = round(total_bank_price + bank_col_6, 2)
                                else:
                                    total_bank_price = round(total_bank_price + bank_col_7, 2)
                        if total_bank_price == total_sage_price:
                            reconciled_bank_row_list1.append(bank_row)
                            cummulative_reconciled_bank_rows.extend(temp_cummulative_reconciled_bank_rows)
                    else:
                        unreconciled_bank_row_list1.append(bank_row)
                elif bank_price_cell != sage_price_cell:
                    unreconciled_bank_row_list1.append(bank_row)

for item in cummulative_reconciled_bank_rows:
    reconciled_bank_row_list1.append(item)

for item in unreconciled_bank_row_list1:
    if item not in reconciled_bank_row_list1:
        final_unreconciled_bank_list.append(item)

final_unreconciled_bank_list = set(final_unreconciled_bank_list)

for bank_row in final_unreconciled_bank_list:
    for col in range(1, 8):
        unposted_bank_sheet.cell(unposted_bank_row, column=col).value = bank_sheet.cell(bank_row, column=col).value
    unposted_bank_row += 1

output_file = "Finished Bank Rec.xlsx"
wb.save(output_file)
print(f"Done. Results saved to {output_file}")
print(f"Reconciled rows   : {sorted(set(reconciled_bank_row_list1))}")
print(f"Unreconciled rows : {sorted(final_unreconciled_bank_list)}")
